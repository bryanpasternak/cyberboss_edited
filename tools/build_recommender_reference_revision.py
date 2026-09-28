from __future__ import annotations

import os
import shutil
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


SOURCE = Path(
    r"C:\Users\19670\OneDrive\文档\chatgpt\推荐人材料\2027Fall_推荐人参考信息_李冠彬老师版_推荐信素材增强版.docx"
)
OUTPUT = Path(
    r"C:\Users\19670\cyberboss\artifacts\recommendation-letter-materials\2027Fall_推荐人参考信息_李冠彬老师版_主题整理修订版.docx"
)

BLACK = "000000"
BODY = "222222"
SECONDARY = "555555"
LIGHT_BLUE = "D9E8F5"
GRID = "D9D9D9"


def set_run_font(
    run,
    *,
    size: float,
    bold: bool = False,
    color: str = BODY,
    latin: str = "Arial",
    east_asia: str = "Microsoft YaHei",
):
    run.font.name = latin
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:eastAsia"), east_asia)


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_table_borders(table, color: str = GRID, size: str = "6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def remove_paragraph(paragraph):
    element = paragraph._element
    element.getparent().remove(element)
    paragraph._p = paragraph._element = None


def format_paragraph(
    paragraph,
    *,
    size: float = 10.5,
    color: str = BODY,
    before: float = 0,
    after: float = 4,
    line_spacing: float = 1.15,
    keep_with_next: bool = False,
    keep_together: bool = True,
):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line_spacing
    fmt.keep_with_next = keep_with_next
    fmt.keep_together = keep_together
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in paragraph.runs:
        set_run_font(run, size=size, color=color, bold=bool(run.bold))


def add_section_heading(doc: Document, text: str):
    paragraph = doc.add_paragraph()
    paragraph.style = doc.styles["Heading 1"]
    run = paragraph.add_run(text)
    set_run_font(run, size=12.0, bold=True, color=BLACK)
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(10)
    fmt.space_after = Pt(5)
    fmt.line_spacing = 1.0
    fmt.keep_with_next = True
    return paragraph


def add_theme_heading(doc: Document, text: str):
    paragraph = doc.add_paragraph()
    paragraph.style = doc.styles["Heading 2"]
    run = paragraph.add_run(text)
    set_run_font(run, size=10.8, bold=True, color=BLACK)
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(7)
    fmt.space_after = Pt(3)
    fmt.line_spacing = 1.0
    fmt.keep_with_next = True
    return paragraph


def add_body(doc: Document, text: str, *, color: str = BODY, after: float = 4):
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(text)
    set_run_font(run, size=10.5, color=color)
    format_paragraph(paragraph, size=10.5, color=color, after=after)
    return paragraph


def add_lead_paragraph(doc: Document, lead: str, text: str):
    paragraph = doc.add_paragraph()
    lead_run = paragraph.add_run(lead)
    set_run_font(lead_run, size=10.5, bold=True, color=BODY)
    body_run = paragraph.add_run(text)
    set_run_font(body_run, size=10.5, color=BODY)
    format_paragraph(paragraph, size=10.5, color=BODY, after=4)
    return paragraph


def add_bullet(doc: Document, lead: str, text: str):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.18)
    paragraph.paragraph_format.first_line_indent = Inches(-0.16)
    bullet = paragraph.add_run("• ")
    set_run_font(bullet, size=10.3, color=BODY)
    lead_run = paragraph.add_run(lead)
    set_run_font(lead_run, size=10.3, bold=True, color=BODY)
    body_run = paragraph.add_run(text)
    set_run_font(body_run, size=10.3, color=BODY)
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(4)
    fmt.line_spacing = 1.12
    fmt.keep_together = True
    return paragraph


def main():
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SOURCE, OUTPUT)
    doc = Document(OUTPUT)

    # Keep the original title, information table, application heading, and application paragraph.
    for paragraph in list(doc.paragraphs)[3:]:
        remove_paragraph(paragraph)

    section = doc.sections[0]
    section.top_margin = Inches(0.68)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)

    title = doc.paragraphs[0]
    title.style = doc.styles["Title"]
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(12)
    title.paragraph_format.keep_with_next = True
    for run in title.runs:
        set_run_font(run, size=18, bold=True, color=BLACK)

    application_heading = doc.paragraphs[1]
    application_heading.style = doc.styles["Heading 1"]
    application_heading.paragraph_format.space_before = Pt(8)
    application_heading.paragraph_format.space_after = Pt(4)
    application_heading.paragraph_format.keep_with_next = True
    for run in application_heading.runs:
        set_run_font(run, size=12, bold=True, color=BLACK)

    application_body = doc.paragraphs[2]
    format_paragraph(application_body, size=10.5, color=BODY, after=5)

    table = doc.tables[0]
    set_table_borders(table)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_shading(cell, LIGHT_BLUE if index in (0, 2) else "FFFFFF")
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
                for run in paragraph.runs:
                    set_run_font(
                        run,
                        size=10.5,
                        bold=index in (0, 2),
                        color=BLACK,
                    )

    add_section_heading(doc, "《模式识别》课程表现与能力证据")

    metadata = doc.add_paragraph()
    r = metadata.add_run("课程成绩：92/100")
    set_run_font(r, size=10.5, bold=True, color=BLACK)
    r = metadata.add_run("    |    修读时间：Spring 2026 (Mar-Jul 2026)")
    set_run_font(r, size=10.5, color=BODY)
    format_paragraph(metadata, size=10.5, color=BODY, after=4)

    add_body(
        doc,
        "课程期间保持全勤，完成全部课程实验及期末作业。以下根据已经完成的实验和报告，整理三项较集中的课程表现，供老师了解和选用；具体评价及推荐信表述以老师的判断为准。",
        after=5,
    )

    add_theme_heading(doc, "1 注重方法原理与完整实现过程")
    add_lead_paragraph(
        doc,
        "PCA 图像压缩实验。",
        "我没有直接调用现成的 PCA 接口，而是自行完成协方差矩阵构造、SVD 分解、主成分投影与图像重建，并应用于 Eigenfaces 与 RGB 图像的降维、压缩与重建。我进一步比较了 10、50、100 和 150 个主成分下的重建效果，以理解主成分数量、信息保留程度与压缩效果之间的关系。",
    )

    add_theme_heading(doc, "2 面对实验偏差时继续追踪原因并调整实现")
    add_lead_paragraph(
        doc,
        "全景图拼接实验。",
        "我实现了 Harris 角点检测，并结合 SIFT、HOG、匹配过滤和 RANSAC 完成双图及四图拼接。针对假阳性角点和低质量匹配对结果造成的影响，我继续进行特征筛选和参数调整，并比较不同描述子在视角变化场景中的实际表现。",
    )
    add_lead_paragraph(
        doc,
        "FixMatch 半监督图像分类实验。",
        "我阅读原论文后，基于 PyTorch 完成了从数据划分、弱强数据增强、伪标签与置信度筛选到联合损失、训练和评估的实现，并使用 USB 标准实现作为参照。当自行实现的结果与参考实现存在性能差异时，我继续从训练步数、学习率调度、指数移动平均和数据增强策略等方面分析可能原因，而没有只停留在记录最终准确率。",
    )

    theme_three = add_theme_heading(doc, "3 主动学习新方向并形成结构化理解")
    theme_three.paragraph_format.page_break_before = True
    add_lead_paragraph(
        doc,
        "期末综述。",
        "我选择视觉词元作为主题，完成了 21 页的《视觉词元的表征与学习：从图像离散化到多模态统一建模》，并自行绘制了 3 张技术示意图。报告按照基础范式、多尺度建模、效率优化和多模态统一的结构，梳理了 ViT、多尺度视觉 Transformer、Token 剪枝与多模态视觉 Token 等方向。",
    )
    add_body(
        doc,
        "在文献整理的基础上，我进一步比较了 Token 压缩中的效率与信息保留、连续与离散视觉表征，以及动态 Token 分配等技术问题，形成了对不同研究路线及其差异的理解。",
    )

    closing = add_body(
        doc,
        "以上内容均对应已经完成的课程实验或期末报告。详细作业、实验结果与报告原文可按老师需要进一步提供。",
        color=SECONDARY,
        after=5,
    )
    for run in closing.runs:
        set_run_font(run, size=9.5, color=SECONDARY)

    add_section_heading(doc, "相关学术与项目背景")
    add_bullet(
        doc,
        "Biomedical RAG 研究 (2025.07-2025.11)：",
        "在 Carnegie Mellon University School of Computer Science 的 Prof. David Woodruff 指导下参与 Biomedical Question Answering / RAG 研究，负责系统代码实现、实验设计与执行、结果分析及论文写作。相关论文已获 ICCSIT 2026 录用，本人为共同一作。",
    )
    add_bullet(
        doc,
        "Multi-channel Proactive AI Companion (2026.03-2026.08)：",
        "参与多通道主动式 AI Agent 系统设计与研发，使用 JavaScript / Node.js 完成多平台接入、跨平台身份与会话管理、长期记忆和主动消息路由等核心模块。",
    )

    add_section_heading(doc, "随附材料")
    add_body(
        doc,
        "个人 CV、成绩单、学校出具的 4.0 制 GPA 证明。如需要申请项目清单、推荐信初稿、课程作业原文或其他补充材料，我会及时整理提供。",
        color=SECONDARY,
        after=0,
    )

    doc.core_properties.title = "2027 Fall 硕士申请推荐人参考信息"
    doc.core_properties.subject = "李冠彬老师推荐信参考素材"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
